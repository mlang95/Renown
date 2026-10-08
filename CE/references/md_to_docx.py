#!/usr/bin/env python3
"""md_to_docx.py — render a prose Markdown rules doc to a styled .docx, injecting
data tables generated live from renown_data wherever a marker appears.

Single source of truth split:
  - prose          -> the .md (RULES_reorganized.md)
  - data tables    -> renown_data.py, via docx_tables.{{TABLE:name}} / {{GLOSSARY}} / {{GLOSSARY:Category}}
  - inline subs    -> {{DEF:term}}, {{TERM:term}}, {{IDX:term}}, {{VERSION}}
  - book furniture -> {{TOC}} (contents), {{INDEX}} (every glossary/IDX term with page numbers),
                      "Renown vX · Page X of Y" footer on every section

{{TERM:x}} prints "**x** — definition" inline; {{IDX:x}} is an invisible index entry for a term the
prose defines itself. Every {{GLOSSARY:Category}} row and every TERM/IDX is an XE field, so the
{{INDEX}} page numbers are built by Word (or LibreOffice) when fields update on open.

This replaces the build_docs.py + Rules_authored.docx path: the prose now lives
in the same .md the wiki consumes, so docx and wiki cannot diverge in wording,
and every embedded table regenerates from canon.

Usage:  python md_to_docx.py RULES_reorganized.md Rules.docx
"""
import re, sys
from docx import Document
from docx.shared import Pt, Twips
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn
import docx_tables as dt
try:
    import renown_data as _rd
    VERSION = str(getattr(_rd, "VERSION", ""))
except Exception:
    VERSION = ""

FONT = "EB Garamond"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HEAD_SIZES = {1: 20, 2: 18, 3: 16, 4: 14, 5: 12, 6: 11}   # matches reference Rules.docx
HEAD_BEFORE = {1: 12, 2: 12, 3: 8, 4: 6, 5: 6, 6: 6}

TABLE_MK    = re.compile(r"^\s*\{\{TABLE:([a-z_]+)\}\}\s*$")
GLOSSARY_MK = re.compile(r"^\s*\{\{GLOSSARY\}\}\s*$")
GLOSSCAT_MK = re.compile(r"^\s*\{\{GLOSSARY:([^}]+)\}\}\s*$")
TOC_MK      = re.compile(r"^\s*\{\{TOC\}\}\s*$")
INDEX_MK    = re.compile(r"^\s*\{\{INDEX\}\}\s*$")
TERM_MK     = re.compile(r"\{\{TERM:([^}]+)\}\}")
IDX_MK      = re.compile(r"\{\{IDX:([^}]+)\}\}")
ACTIONS_MK  = re.compile(r"^\s*\{\{ACTIONS:([A-Za-z]+)\}\}\s*$")
LIST_MK     = re.compile(r"^\s*\{\{LIST:([A-Z_]+)\}\}\s*$")
COLS_MK     = re.compile(r"^\s*\{\{COLS:(\d)\}\}\s*$")
DEF_MK      = re.compile(r"\{\{DEF:([^}]+)\}\}")
VAL_MK      = re.compile(r"\{\{VAL:([^}]+)\}\}")
VERSION_MK  = re.compile(r"\{\{VERSION\}\}")
INLINE      = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*|`.+?`)")


def _subs(text):
    """Inline {{VERSION}} / {{VAL:path}} / {{DEF:term}} substitutions (run before run parsing)."""
    text = VERSION_MK.sub(VERSION, text)
    text = VAL_MK.sub(lambda m: dt.value(m.group(1).strip()), text)
    text = DEF_MK.sub(lambda m: dt.definition(m.group(1).strip()) or m.group(0), text)
    text = TERM_MK.sub(lambda m: _term(m.group(1).strip()), text)
    text = IDX_MK.sub(lambda m: dt.xe_mark(m.group(1).strip()), text)
    # display layer (NAME_DISPLAY): rename what players read; markers already resolved above
    return getattr(dt.rd, "display_md", lambda x: x)(text)


def _term(term):
    key, d = dt.lookup(term)
    if key is None:
        raise KeyError(f"TERM '{term}' is not in GLOSSARY")
    return f"**{term}** — {d}" + dt.xe_mark(key)


def _runs(paragraph, text, base_bold=False, base_size=None):
    """Add runs to a paragraph, parsing **bold**, *italic*, ***both***, `code`; XE sentinels -> index fields."""
    text = _subs(text)
    chunks = dt.XE_RE.split(text)          # [text, term, text, term, ...]
    for j, chunk in enumerate(chunks):
        if j % 2:
            for r in parse_xml(f'<w:root xmlns:w="{W_NS}">{dt.xe_field_xml(chunk)}</w:root>'):
                paragraph._p.append(r)
        else:
            _plain_runs(paragraph, chunk, base_bold, base_size)


def _plain_runs(paragraph, text, base_bold=False, base_size=None):
    for part in INLINE.split(text):
        if part == "":
            continue
        bold, ital, body = base_bold, False, part
        if part.startswith("***") and part.endswith("***"):
            bold, ital, body = True, True, part[3:-3]
        elif part.startswith("**") and part.endswith("**"):
            bold, body = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*"):
            ital, body = True, part[1:-1]
        elif part.startswith("`") and part.endswith("`"):
            body = part[1:-1]
        r = paragraph.add_run(body)
        r.font.name = FONT
        if base_size:
            r.font.size = Pt(base_size)
        r.bold = bold
        r.italic = ital


def _colbreak(n):
    """A continuous section break. Per OOXML, a sectPr in a paragraph defines the
    section ENDING at that paragraph, so {{COLS:N}} sets N columns for everything
    since the previous break up to here. The final (body) sectPr stays 1-column."""
    cols = ('<w:cols w:space="708"/>' if n == 1
            else f'<w:cols w:num="{n}" w:space="360" w:equalWidth="1"/>')
    return ('<w:p><w:pPr><w:sectPr><w:type w:val="continuous"/>'
            '<w:pgSz w:w="11906" w:h="16838"/>'
            '<w:pgMar w:top="720" w:right="720" w:bottom="720" w:left="720" '
            'w:header="708" w:footer="708" w:gutter="0"/>'
            f'{cols}</w:sectPr></w:pPr></w:p>')


def _inject(doc, ooxml):
    """Insert one or more OOXML sibling elements (a <w:tbl> or several <w:p>) at the
    current end of body content — before the trailing <w:sectPr>, where add_paragraph
    also inserts — so injected tables land at their marker position, not document end."""
    root = parse_xml(f'<w:root xmlns:w="{W_NS}">{ooxml}</w:root>')
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    for child in list(root):
        if sectPr is not None:
            sectPr.addprevious(child)
        else:
            body.append(child)


def _field_par(instr, placeholder, size=20):
    rpr = f'<w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}"/><w:sz w:val="{size}"/></w:rPr>'
    return (f'<w:p><w:r>{rpr}<w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
            f'<w:r>{rpr}<w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
            f'<w:r>{rpr}<w:fldChar w:fldCharType="separate"/></w:r>'
            f'<w:r>{rpr}<w:t xml:space="preserve">{placeholder}</w:t></w:r>'
            f'<w:r>{rpr}<w:fldChar w:fldCharType="end"/></w:r></w:p>')


def _toc():
    return (_field_par('TOC \\o "1-2" \\h \\z \\u', "Update fields (F9) to build the table of contents.")
            + '<w:p><w:r><w:br w:type="page"/></w:r></w:p>')


def _index():
    # \c 2 = two columns, \h "A" = letter group headings
    return _field_par('INDEX \\c "2" \\h "A" \\e ", " \\z "1033"', "Update fields (F9) to build the index.")


def _footers(doc):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        xml = (f'<w:r><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}"/><w:sz w:val="16"/></w:rPr>'
               f'<w:t xml:space="preserve">Renown v{VERSION}   \u00b7   Page </w:t></w:r>')
        def fld(instr):
            rpr = f'<w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}"/><w:sz w:val="16"/></w:rPr>'
            return (f'<w:r>{rpr}<w:fldChar w:fldCharType="begin"/></w:r><w:r>{rpr}<w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
                    f'<w:r>{rpr}<w:fldChar w:fldCharType="separate"/></w:r><w:r>{rpr}<w:t>1</w:t></w:r><w:r>{rpr}<w:fldChar w:fldCharType="end"/></w:r>')
        xml += fld("PAGE") + (f'<w:r><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}"/><w:sz w:val="16"/></w:rPr>'
                              f'<w:t xml:space="preserve"> of </w:t></w:r>') + fld("NUMPAGES")
        for r in parse_xml(f'<w:root xmlns:w="{W_NS}">{xml}</w:root>'):
            p._p.append(r)


def _index_styles(doc):
    """Index 1 / Index Heading: compact 9pt EB Garamond. No tab stop: Word renders 'Term, 12' (\\e ", "),
    LibreOffice applies its own dotted right tab at the column edge."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Inches
    names = {st.name for st in doc.styles}
    for name, bold, size in (("Index 1", False, 9), ("Index Heading", True, 10)):
        st = doc.styles[name] if name in names else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if qn("w:customStyle") in st.element.attrib:
            del st.element.attrib[qn("w:customStyle")]
        st.font.name = FONT; st.font.size = Pt(size); st.font.bold = bold
        pf = st.paragraph_format
        pf.space_before = Pt(4 if bold else 0); pf.space_after = Pt(0); pf.line_spacing = 1.0
        if not bold:
            pf.left_indent = Inches(0.15); pf.first_line_indent = Inches(-0.15)


def _update_fields_on_open(doc):
    st = doc.settings.element
    if st.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); st.append(uf)


def _gfm_cells(row):
    row = row.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [c.strip() for c in row.split("|")]


def render(md_path, out_path):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Twips(11906); sec.page_height = Twips(16838)   # A4
    sec.top_margin = sec.bottom_margin = Twips(720)                # 0.5"
    sec.left_margin = sec.right_margin = Twips(720)
    sec.header_distance = Twips(708); sec.footer_distance = Twips(708)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0

    i, n = 0, len(lines)
    list_style = {"-": "List Bullet", "*": "List Bullet"}
    while i < n:
        raw = lines[i]
        ln = raw.rstrip()
        s = ln.strip()
        if not s:
            i += 1
            continue

        # block markers
        m = TABLE_MK.match(ln)
        if m:
            _inject(doc, dt.get(m.group(1)))
            i += 1
            continue
        gm = GLOSSCAT_MK.match(ln)
        if gm:
            _inject(doc, dt.glossary_category(gm.group(1).strip()))
            i += 1
            continue
        if TOC_MK.match(ln):
            _inject(doc, _toc())
            i += 1
            continue
        if INDEX_MK.match(ln):
            _inject(doc, _index())
            i += 1
            continue
        if GLOSSARY_MK.match(ln):
            _inject(doc, dt.glossary_block())
            i += 1
            continue
        am = ACTIONS_MK.match(ln)
        if am:
            _inject(doc, dt.actions(am.group(1)))
            i += 1
            continue
        lm0 = LIST_MK.match(ln)
        if lm0:
            _inject(doc, dt.list_block(lm0.group(1)))
            i += 1
            continue
        cm = COLS_MK.match(ln)
        if cm:
            _inject(doc, _colbreak(int(cm.group(1))))
            i += 1
            continue

        # heading
        hm = re.match(r"^(#{1,6})\s+(.*)$", s)
        if hm:
            lvl = len(hm.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(HEAD_BEFORE.get(lvl, 6))
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            _runs(p, hm.group(2), base_bold=True, base_size=HEAD_SIZES.get(lvl, 11))
            ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(lvl - 1))
            p._p.get_or_add_pPr().append(ol)
            i += 1
            continue

        # GFM pipe table: header row, separator (|---|), then body rows
        if "|" in s and i + 1 < n and re.match(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$", lines[i + 1].strip()) and "-" in lines[i + 1]:
            header = _gfm_cells(s)
            i += 2
            body = []
            while i < n and "|" in lines[i] and lines[i].strip():
                body.append(_gfm_cells(lines[i]))
                i += 1
            # resolve {{VAL:}} / {{VERSION}} / {{DEF:}} inside cells too (they bypass _runs)
            header = [_subs(c) for c in header]
            body = [[_subs(c) for c in r] for r in body]
            _inject(doc, dt._table(header, body))   # same dense style as data tables
            continue

        # list item
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", raw)
        if lm:
            marker = lm.group(2)
            if marker[0].isdigit():
                # Keep the source's own numbering. Word's "List Number" style uses
                # one document-wide counter, so successive lists would read 34,35…
                # instead of restarting at 1; emit the authored number as text.
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.first_line_indent = Pt(-18)
                _runs(p, f"{marker} {lm.group(3)}")
            else:
                p = doc.add_paragraph(style="List Bullet")
                _runs(p, lm.group(3))
            i += 1
            continue

        # paragraph
        p = doc.add_paragraph()
        _runs(p, s)
        i += 1

    _footers(doc)
    _index_styles(doc)
    _update_fields_on_open(doc)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit("usage: python md_to_docx.py RULES_reorganized.md Rules.docx")
    out = render(sys.argv[1], sys.argv[2])
    print("wrote " + out)